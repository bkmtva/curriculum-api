import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import json
file_path = "src/excel_files/template.xlsx"
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font
import pandas as pd


def to_excel(dict_train):
    # with open(response_file_path, encoding='utf8') as response_file:
    #     dict_train = json.load(response_file)
    dict_train = json.loads(dict_train)
    all_courses = dict_train.get('courses')
    formatted_data = []
    for course_data in all_courses:
        course_info = course_data['course']
        subcourses = course_info.get('subcourses')

        # If there are subcourses, iterate over them
        if subcourses:
            for subcourse in subcourses:
                formatted_data.append({
                    'semester': course_data['semester'],
                    'course_id': course_data['course_id'],
                    'order_in_semester': course_data['order_in_semester'],
                    'title': subcourse['title'],
                    'title_kz': subcourse.get('title_kz', course_info.get('title_kz', '')),
                    # Use course title_kz if subcourse title_kz is missing
                    'title_ru': subcourse.get('title_ru', course_info.get('title_ru', '')),
                    # Use course title_ru if subcourse title_ru is missing
                    'teor': subcourse['teor'],
                    'cr': subcourse['cr'],
                    'term': subcourse['term'],
                    'course_code': course_info['course_code'],
                    'pr': subcourse['pr'],
                    'ects': subcourse['ects'],
                    'subcourse_title': None  # Set to None for subcourses
                })
        else:
            formatted_data.append({
                'semester': course_data['semester'],
                'course_id': course_data['course_id'],
                'order_in_semester': course_data['order_in_semester'],
                'title': course_info['title'],
                'title_kz': course_info.get('title_kz'),
                'title_ru': course_info.get('title_ru'),
                'teor': course_info['teor'],
                'cr': course_info['cr'],
                'term': course_info['term'],
                'course_code': course_info['course_code'],
                'pr': course_info['pr'],
                'ects': course_info['ects'],
                'subcourse_title': None  # Set to None for main courses without subcourses
            })

    # Creating DataFrame
    df = pd.DataFrame(formatted_data)
    df['combined_titles'] = df['title_kz'] + ' / ' + df['title_ru'] + ' / ' + df['title']
    df['pr_value'] = df['pr'].apply(lambda x: sum(map(int, x.split('+'))))
    save_file(file_path, df, dict_train)
    return df


def save_file(file_path, df, dict_train):
    wb = load_workbook(filename=file_path)
    ws = wb['CS Curriculum 2023-2024(main)']

    start_col = 4
    start_row = 28
    semesters = df['semester'].unique()

    semesters.sort()

    bold_font = Font(name='Times New Roman', size=12, bold=True)
    regular_font = Font(name='Times New Roman', size=12)

    for semester in semesters:
        df_semester = df[df['semester'] == semester]

        df_semester = df_semester.sort_values(by='order_in_semester')

        merge_range = f'{get_column_letter(start_col)}{start_row}:{get_column_letter(start_col + 4)}{start_row}'
        ws.merge_cells(merge_range)

        ws[f'{get_column_letter(start_col)}{start_row}'].value = f"{semester} семестр"
        ws[f'{get_column_letter(start_col)}{start_row}'].alignment = Alignment(horizontal='center')
        ws[f'{get_column_letter(start_col)}{start_row}'].font = bold_font

        start_row += 1

        cr_start_row = start_row

        ects_start_row = start_row

        prev_order = None
        max_cr = 0
        max_ects = 0

        for r, course in enumerate(df_semester.itertuples(), start=start_row):
            if prev_order == course.order_in_semester and r != start_row:
                ws.merge_cells(f'{get_column_letter(start_col)}{r}:{get_column_letter(start_col)}{r}')
                ws[f'{get_column_letter(start_col)}{r}'].value = None
            else:
                ws[f'{get_column_letter(start_col)}{r}'].value = course.course_code

            ws[f'{get_column_letter(start_col + 1)}{r}'].value = course.combined_titles
            ws[f'{get_column_letter(start_col + 2)}{r}'].value = course.cr
            ws[f'{get_column_letter(start_col + 3)}{r}'].value = course.ects
            ws[f'{get_column_letter(start_col + 4)}{r}'].value = course.pr_value

            if prev_order == course.order_in_semester:
                max_cr = max(max_cr, course.cr)
                max_ects = max(max_ects, course.ects)
            else:
                if prev_order is not None:
                    merge_range_cr = f'{get_column_letter(start_col + 2)}{cr_start_row}:{get_column_letter(start_col + 2)}{r - 1}'
                    ws.merge_cells(merge_range_cr)
                    ws[f'{get_column_letter(start_col + 2)}{cr_start_row}'].value = max_cr

                    merge_range_ects = f'{get_column_letter(start_col + 3)}{ects_start_row}:{get_column_letter(start_col + 3)}{r - 1}'
                    ws.merge_cells(merge_range_ects)
                    ws[f'{get_column_letter(start_col + 3)}{ects_start_row}'].value = max_ects

                cr_start_row = r
                ects_start_row = r
                max_cr = course.cr
                max_ects = course.ects
                prev_order = course.order_in_semester

            for col in range(start_col, start_col + 5):
                ws[f'{get_column_letter(col)}{r}'].font = regular_font

        merge_range_cr = f'{get_column_letter(start_col + 2)}{cr_start_row}:{get_column_letter(start_col + 2)}{start_row + len(df_semester) - 1}'
        ws.merge_cells(merge_range_cr)
        ws[f'{get_column_letter(start_col + 2)}{cr_start_row}'].value = max_cr

        merge_range_ects = f'{get_column_letter(start_col + 3)}{ects_start_row}:{get_column_letter(start_col + 3)}{start_row + len(df_semester) - 1}'
        ws.merge_cells(merge_range_ects)
        ws[f'{get_column_letter(start_col + 3)}{ects_start_row}'].value = max_ects

        for row in ws.iter_rows(min_row=start_row, max_row=start_row + len(df_semester) - 1, min_col=start_col + 1,
                                max_col=start_col + 1):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, horizontal='center')

        ws[f'{get_column_letter(start_col)}{start_row + len(df_semester)}'].value = 'Барлығы/Итого/Total'
        ws.merge_cells(
            f'{get_column_letter(start_col)}{start_row + len(df_semester)}:{get_column_letter(start_col + 1)}{start_row + len(df_semester)}')

        df_semester['cr'] = df_semester['cr'].astype(int)
        df_semester['ects'] = df_semester['ects'].astype(int)

        total_cr = sum(df_semester.groupby(['semester', 'order_in_semester'])['cr'].max())
        total_ects = sum(df_semester.groupby(['semester', 'order_in_semester'])['ects'].max())
        ws[f'{get_column_letter(start_col + 2)}{start_row + len(df_semester)}'].value = total_cr
        ws[f'{get_column_letter(start_col + 3)}{start_row + len(df_semester)}'].value = total_ects

        for col in range(start_col, start_col + 4):
            ws[f'{get_column_letter(col)}{start_row + len(df_semester)}'].font = bold_font
            ws[f'{get_column_letter(col)}{start_row + len(df_semester)}'].alignment = Alignment(horizontal='center')

        start_row += len(df_semester) + 1

    program_in_kz = 'Білім беру бағдарламасы: ' + dict_train.get('program').get('code', '') + ' ' + dict_train.get(
        'program').get('title_kz', '')
    program_in_ru = 'Образовательная программа: ' + dict_train.get('program').get('code', '') + ' ' + dict_train.get(
        'program').get('title_ru', '')
    program_in_en = 'Program: ' + dict_train.get('program').get('code') + ' ' + dict_train.get('program').get('title')

    degree_in_kz = 'Берілетін дәреже: ' + dict_train.get('program').get('code')
    degree_in_ru = 'Присуждаемая степень: ' + dict_train.get('program').get('code')
    degree_in_en = 'Degree: ' + dict_train.get('program').get('code')

    ws.cell(row=17, column=1, value=program_in_kz)
    ws.cell(row=17, column=5, value=program_in_en)
    ws.cell(row=17, column=7, value=program_in_ru)

    ws.cell(row=18, column=1, value=degree_in_kz)
    ws.cell(row=18, column=5, value=degree_in_en)
    ws.cell(row=18, column=7, value=degree_in_ru)
    wb.save("src/excel_files/curriculum.xlsx")

